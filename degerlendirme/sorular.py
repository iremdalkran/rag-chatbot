"""
sorular.py
----------
Soru dosyası (Excel) biçimi: okuma, doğrulama ve boş şablon üretme.

Sütunlar:
  Soru            — sisteme sorulacak soru
  Doğru cevap     — beklenen cevap ("dokümanda olmayan" sorularda boş bırakılabilir)
  Kaynak dosya    — cevabın bulunduğu dosyanın adı (ör. personel_yonetmeligi.pdf, satislar.xlsx)
  Kaynak sayfa    — PDF sayfa numarası: 4, 4-5 veya 4, 6 (bilinmiyorsa ya da Excel'se boş)
  Soru türü       — doküman / Excel / dokümanda olmayan
"""

from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Set

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = ["Soru", "Doğru cevap", "Kaynak dosya", "Kaynak sayfa", "Soru türü"]
TYPE_DOC = "doküman"
TYPE_DATA = "Excel"
TYPE_NONE = "dokümanda olmayan"
TYPES = [TYPE_DOC, TYPE_DATA, TYPE_NONE]

_FOLD = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")


class QuestionFileError(Exception):
    """Soru dosyasındaki, kullanıcıya gösterilecek hatalar."""


@dataclass
class Question:
    row: int                 # Excel'deki satır numarası (rapordan dosyaya dönmek kolay olsun diye)
    text: str
    expected: str
    source_file: str
    source_pages: Optional[Set[int]]
    kind: str                # TYPES içinden biri


def _fold(text: str) -> str:
    return " ".join(str(text).translate(_FOLD).lower().split())


def normalize_type(value) -> Optional[str]:
    folded = _fold(value or "")
    if not folded:
        return None
    if "olmayan" in folded or folded in ("yok", "cevapsiz", "kapsam disi"):
        return TYPE_NONE
    if folded in ("excel", "veri", "tablo", "csv", "sql"):
        return TYPE_DATA
    if folded.startswith("dokuman") or folded in ("pdf", "belge", "word"):
        return TYPE_DOC
    return None


def parse_pages(value) -> Optional[Set[int]]:
    """'4' → {4}, '4-6' → {4, 5, 6}, '4, 7' → {4, 7}, boş → None (sayfa kontrolü yapılmaz)."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return {int(value)}
    pages: Set[int] = set()
    for part in str(value).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, _, end = part.partition("-")
            pages.update(range(int(float(start)), int(float(end)) + 1))
        else:
            pages.add(int(float(part)))
    return pages or None


def load_questions(path: Path) -> List[Question]:
    try:
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    except Exception as e:
        raise QuestionFileError(f"Soru dosyası açılamadı: {path} ({e})") from e

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise QuestionFileError("Soru dosyası boş.")
    header = [_fold(c or "") for c in rows[0]]
    index = {}
    for col in COLUMNS:
        try:
            index[col] = header.index(_fold(col))
        except ValueError:
            raise QuestionFileError(
                f"'{col}' sütunu bulunamadı. İlk satırda şu başlıklar olmalı: {', '.join(COLUMNS)}"
            ) from None

    questions, problems = [], []
    for number, row in enumerate(rows[1:], start=2):
        def cell(col):
            i = index[col]
            value = row[i] if i < len(row) else None
            return "" if value is None else str(value).strip()

        text = cell("Soru")
        if not text:
            continue  # boş satırlar atlanır
        kind = normalize_type(cell("Soru türü"))
        if kind is None:
            problems.append(f"{number}. satır: 'Soru türü' şunlardan biri olmalı: {', '.join(TYPES)}")
            continue
        try:
            pages = parse_pages(row[index["Kaynak sayfa"]] if index["Kaynak sayfa"] < len(row) else None)
        except ValueError:
            problems.append(f"{number}. satır: 'Kaynak sayfa' anlaşılamadı (örnek: 4 veya 4-5 veya 4, 7)")
            continue
        expected = cell("Doğru cevap")
        if kind != TYPE_NONE and not expected:
            problems.append(f"{number}. satır: 'Doğru cevap' boş (sadece 'dokümanda olmayan' sorularda boş olabilir)")
            continue
        questions.append(Question(number, text, expected, cell("Kaynak dosya"), pages, kind))

    if problems:
        raise QuestionFileError("Soru dosyasında düzeltilmesi gereken satırlar var:\n  - " + "\n  - ".join(problems))
    if not questions:
        raise QuestionFileError("Soru dosyasında hiç soru yok.")
    return questions


# --- Şablon ---

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")

INSTRUCTIONS = [
    ("Sütun", "Ne yazılmalı?", "Örnek"),
    ("Soru", "Sisteme sorulacak soru. Kullanıcıların gerçekten soracağı şekilde yazın.",
     "Beş yıldan az çalışan bir personelin yıllık izni kaç gündür?"),
    ("Doğru cevap", "Beklenen cevap. Kısa ve net olsun; önemli olan sayı, ad, tarih gibi bilgiler. "
     "'dokümanda olmayan' sorularda boş bırakın.", "14 iş günü"),
    ("Kaynak dosya", "Cevabın bulunduğu dosyanın adı, uzantısıyla birlikte. Dosya, soru dosyasıyla aynı "
     "klasörde olmalı. 'dokümanda olmayan' sorularda boş bırakın.", "personel_yonetmeligi.pdf"),
    ("Kaynak sayfa", "PDF'te cevabın olduğu sayfa. Birden fazla sayfa için 4-5 veya 4, 7 yazın. "
     "Word/TXT/Excel için ya da bilmiyorsanız boş bırakın.", "2"),
    ("Soru türü", "Listeden seçin: doküman (PDF/Word/TXT'den cevaplanır), Excel (tablodan hesaplanır), "
     "dokümanda olmayan (sistem 'bulamadım' demeli, uydurmamalı).", "doküman"),
    ("", "", ""),
    ("İpuçları", "• Her türden soru ekleyin; özellikle 'dokümanda olmayan' sorular, sistemin uydurup "
     "uydurmadığını ölçer.\n• 20-50 soru iyi bir başlangıçtır.\n• Ayarları değiştirmeden önce ve sonra "
     "aynı soru dosyasıyla çalıştırıp raporları karşılaştırın.", ""),
]


def write_question_file(path: Path, rows: Optional[list] = None) -> None:
    """Boş şablonu (rows=None) ya da verilen satırlarla doldurulmuş bir soru dosyası yazar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sorular"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows or []:
        ws.append(list(row))
    for col, width in zip("ABCDE", (60, 45, 28, 14, 20)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    validation = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=False,
                                showErrorMessage=True, errorTitle="Geçersiz soru türü",
                                error="Listeden seçin: " + ", ".join(TYPES))
    ws.add_data_validation(validation)
    validation.add("E2:E1000")

    guide = wb.create_sheet("Nasıl doldurulur")
    for row in INSTRUCTIONS:
        guide.append(row)
    for cell in guide[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for col, width in zip("ABC", (16, 80, 50)):
        guide.column_dimensions[col].width = width
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
