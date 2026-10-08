"""
karsilastir.py
--------------
Ayar karşılaştırması: değerlendirmeyi her seferinde TEK bir ayarı değiştirerek tekrar tekrar
çalıştırır ve sonuçları tek bir karşılaştırma tablosunda toplar.

Denenen ayarlar (varsayılanlar, komut satırından değiştirilebilir):
  - Parça boyutu (CHUNK_CHARS):        600, 1200, 2000 karakter
  - Bulunan parça sayısı (TOP_K):       3, 6, 10
  - Cevap modeli (CHAT_MODEL):          qwen3:8b, qwen3:14b
  - Doküman arama yöntemi (RAG_METHOD): istenirse vector ve pageindex (--yontemler vector,pageindex)

"Temel" ayar, .env'deki (ya da varsayılan) mevcut ayardır. Her deneme temelden yalnızca bir
ayarla ayrılır; böylece bir fark çıkarsa onun hangi ayardan geldiği bellidir. Sonunda her ayar
için en iyi değer seçilir; bu değerlerin birleşimi daha önce denenmemişse o da ayrıca çalıştırılıp
doğrulanır (ayarlar birbirini etkileyebilir).

Bütün denemelerde AYNI hakem modeli kullanılır; aksi halde puanlar karşılaştırılamaz.

Kullanım:  bash ayar_karsilastir.sh ornekler/degerlendirme/sorular.xlsx
"""

import argparse
import os
import shutil
import statistics
import sys
import tempfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_JUDGE = "qwen3:30b-a3b"
# Yaklaşık indirme boyutları (GB) — sadece disk yeri uyarısı için.
MODEL_SIZES_GB = {"qwen3:4b": 2.6, "qwen3:8b": 5.2, "qwen3:14b": 9.3, "qwen3:30b-a3b": 18.6, "bge-m3": 1.2}


def _floats(text: str) -> list:
    return [int(x) for x in text.split(",") if x.strip()]


def _num(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


METHOD_LABEL = {"vector": "Vektör", "pageindex": "PageIndex"}


def _variants(base: dict, chunk_sizes: list, top_ks: list, models: list, methods: list = ()) -> list:
    """Temel + her ayar için temelden tek farkla ayrılan denemeler."""
    variants = [{"name": "temel", "factor": None, "change": "— (temel ayar)", **base}]
    for size in chunk_sizes:
        if size != base["chunk"]:
            variants.append({"name": f"parca-{size}", "factor": "chunk",
                             "change": f"Parça boyutu {base['chunk']} → {size}", **{**base, "chunk": size}})
    for k in top_ks:
        if k != base["top_k"]:
            variants.append({"name": f"topk-{k}", "factor": "top_k",
                             "change": f"Bulunan parça sayısı {base['top_k']} → {k}", **{**base, "top_k": k}})
    for model in models:
        if model != base["model"]:
            variants.append({"name": "model-" + model.replace(":", "-"), "factor": "model",
                             "change": f"Cevap modeli {base['model']} → {model}", **{**base, "model": model}})
    for method in methods:
        if method != base["method"]:
            variants.append({"name": f"yontem-{method}", "factor": "method",
                             "change": f"Arama yöntemi {METHOD_LABEL[base['method']]} → {METHOD_LABEL[method]}",
                             **{**base, "method": method}})
    return variants


def _apply(variant: dict, temp_root: Path) -> None:
    """Ayarları uygular ve her deneme için temiz, boş bir veri klasörü açar
    (parça boyutu değişince dokümanların yeniden işlenmesi gerekir)."""
    from app import config, db, documents

    config.CHAT_MODEL = variant["model"]
    config.CHUNK_CHARS = variant["chunk"]
    config.RETRIEVAL_TOP_K = variant["top_k"]
    config.RAG_METHOD = variant["method"]
    data_dir = Path(tempfile.mkdtemp(prefix=variant["name"] + "-", dir=temp_root))
    config.DATA_DIR = data_dir
    config.DB_PATH = data_dir / "app.db"
    config.TABLES_DIR = data_dir / "tables"
    db.init_db()
    with documents._cache_lock:
        documents._cache.clear()


def _row(variant: dict, outcome: dict, repeats: list) -> dict:
    """Bir denemenin (gerekirse tekrarlarının ortalaması alınmış) sonuç satırı."""
    def mean(key, by_type=None):
        values = []
        for summary in repeats:
            source = summary["by_type"].get(by_type, {}) if by_type else summary
            value = source.get(key)
            value = _num(value) if not isinstance(value, float) else value
            if value is not None:
                values.append(value)
        return statistics.mean(values) if values else None

    def check(key, kind):
        values = [s["by_type"][kind][key] for s in repeats if kind in s["by_type"] and s["by_type"][kind][key] is not None]
        return statistics.mean(values) if values else None

    return {
        "name": variant["name"], "change": variant["change"], "factor": variant["factor"],
        "model": variant["model"], "chunk": variant["chunk"], "top_k": variant["top_k"],
        "method": variant["method"], "chunk_count": outcome["chunk_count"],
        "prep": statistics.mean(outcome["prep"]) if outcome.get("prep") else None,
        "node_count": outcome.get("node_count", 0),
        "score": mean("score"),
        "score_doc": mean("score", "doküman"), "score_data": mean("score", "Excel"),
        "score_none": mean("score", "dokümanda olmayan"),
        "retrieved": check("retrieved", "doküman"), "cited": check("cited", "doküman"),
        "numbers": statistics.mean([v for v in (check("numbers", "doküman"), check("numbers", "Excel")) if v is not None])
        if any(check("numbers", k) is not None for k in ("doküman", "Excel")) else None,
        "wait": mean("avg_wait"), "first": mean("avg_first"), "total": mean("avg_seconds"),
        "reports": [str(o) for o in outcome["reports"]],
    }


def _pick(candidates: list, tolerance: float) -> dict:
    """En yüksek doğruluğa bir soru mesafesinde olanlar arasından algılanan beklemesi en kısa olanı seçer.
    (İkili karşılaştırma yerine bu yöntem: sonuç denemelerin sırasına bağlı olmaz.)"""
    scored = [c for c in candidates if c["score"] is not None]
    if not scored:
        return candidates[0]
    best_score = max(c["score"] for c in scored)
    close = [c for c in scored if c["score"] >= best_score - tolerance]
    return min(close, key=lambda c: c["wait"] if c["wait"] is not None else float("inf"))


def recommend(rows: list, question_count: int) -> dict:
    """Her ayar için en iyi değeri, temel + o ayarın denemeleri arasından seçer."""
    tolerance = 1.0 / max(question_count, 1) + 1e-9   # en fazla bir soruluk fark "aynı" sayılır
    base = next(r for r in rows if r["name"] == "temel")
    best = {factor: _pick([base] + [r for r in rows if r["factor"] == factor], tolerance)
            for factor in ("chunk", "top_k", "model", "method")}
    return {
        "model": best["model"]["model"], "chunk": best["chunk"]["chunk"], "top_k": best["top_k"]["top_k"],
        "method": best["method"]["method"],
        "winners": best, "tolerance": tolerance,
    }


def _pct(value):
    return "—" if value is None else f"%{value * 100:.0f}"


def _sec(value):
    return "—" if value is None else f"{value:.1f}"


def write_comparison(path: Path, rows: list, choice: dict, judge: str, question_file: Path, repeats: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill, header_font = PatternFill("solid", fgColor="4F46E5"), Font(bold=True, color="FFFFFF")
    best_fill = PatternFill("solid", fgColor="D1FAE5")
    wb = Workbook()
    ws = wb.active
    ws.title = "Karşılaştırma"
    ws["A1"] = "Ayar Karşılaştırması"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"{datetime.now():%d.%m.%Y %H:%M} · soru dosyası: {question_file.name} · hakem: {judge} · "
                f"her deneme {repeats} kez çalıştırıldı")
    columns = [("Deneme", 22), ("Değişen ayar", 30), ("Yöntem", 11), ("Cevap modeli", 13), ("Parça boyutu", 9),
               ("Top-k", 7), ("Parça sayısı", 9), ("Bölüm sayısı (PageIndex)", 10), ("Hazırlık (sn)", 10), ("Doğruluk", 10), ("Doküman", 10), ("Excel", 9), ("Dokümanda olmayan", 11),
               ("Aramada bulundu", 11), ("Cevapta gösterildi", 11), ("Sayılar tuttu", 10),
               ("Algılanan bekleme (sn)", 11), ("İlk kelime (sn)", 10), ("Ort. toplam süre (sn)", 11),
               ("Doğruluk farkı (temele göre)", 12), ("Bekleme farkı (sn)", 11), ("Rapor(lar)", 50)]
    for col, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=4, column=col, value=name)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    base = next(r for r in rows if r["name"] == "temel")
    winners = {id(r) for r in choice["winners"].values()}
    for i, r in enumerate(rows, start=5):
        d_score = None if r["score"] is None or base["score"] is None else (r["score"] - base["score"]) * 100
        d_wait = None if r["wait"] is None or base["wait"] is None else r["wait"] - base["wait"]
        values = [r["name"], r["change"], METHOD_LABEL[r["method"]], r["model"], r["chunk"], r["top_k"],
                  r["chunk_count"], r["node_count"] or "—", _sec(r["prep"]), _pct(r["score"]),
                  _pct(r["score_doc"]), _pct(r["score_data"]), _pct(r["score_none"]), _pct(r["retrieved"]),
                  _pct(r["cited"]), _pct(r["numbers"]), _sec(r["wait"]), _sec(r["first"]), _sec(r["total"]),
                  "—" if d_score is None or r is base else f"{d_score:+.0f} puan",
                  "—" if d_wait is None or r is base else f"{d_wait:+.1f}", "\n".join(Path(p).name for p in r["reports"])]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r["name"] == "onerilen" or (id(r) in winners and r is not base):
                cell.fill = best_fill
    ws.freeze_panes = "B5"

    notes_row = 6 + len(rows)
    notes = [
        ("Önerilen ayar", f"RAG_METHOD={choice['method']}  ·  CHAT_MODEL={choice['model']}  ·  "
                          f"CHUNK_CHARS={choice['chunk']}  ·  RETRIEVAL_TOP_K={choice['top_k']}"),
        ("Nasıl seçildi?", "Her ayar için, o ayarın denendiği satırlar ile temel arasından: önce doğruluk; doğruluk farkı "
                           "en fazla bir soruysa (gürültü sayılır) algılanan beklemesi en kısa olan. Yeşil satırlar "
                           "seçilen değerlerdir."),
        ("Algılanan bekleme", "Kullanıcının ekranda bir şey görene kadar beklediği süre (doküman sorusunda ilk kelime, "
                              "Excel sorusunda cevabın tamamı)."),
        ("Dikkat", "Küçük bir soru setinde bir sorunun doğru ya da yanlış çıkması puanı birkaç puan oynatır. Bu yüzden "
                   "bir soruluk farklar 'aynı' sayıldı. Sonuç ancak soru seti kadar güvenilirdir; gerçek "
                   "dokümanlarınızdan hazırlanmış bir setle tekrar çalıştırmanız önerilir."),
        ("Hazırlık", "Dokümanların yüklenip aranabilir hale gelmesi için geçen süre (bir kez ödenir). PageIndex'te "
                     "buna içindekiler ağacının kurulması (başlık çıkarma + bölüm özetleri) da dahildir."),
        ("Aramada bulundu", "Vektörde: bulunan parçalar arasında doğru sayfa var mı? PageIndex'te: modelin seçip okuduğu "
                            "sayfalar arasında doğru sayfa var mı?"),
        ("Parça boyutu değişirse", "Mevcut dokümanların silinip yeniden yüklenmesi gerekir (parçalar yükleme anında oluşur)."),
    ]
    for i, (label, text) in enumerate(notes):
        ws.cell(row=notes_row + i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=notes_row + i, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=notes_row + i, start_column=2, end_row=notes_row + i, end_column=12)
        ws.row_dimensions[notes_row + i].height = 32
    wb.save(path)


def run_comparison(question_file: Path, files_dir: Path, output_dir: Path, judge: str, chunk_sizes: list,
                   top_ks: list, models: list, repeats: int = 1, limit: int = None, log=print,
                   pull_missing: bool = True, methods: list = ()) -> dict:
    from app import config, llm

    from .calistir import EvaluationError, run_evaluation
    from .sorular import load_questions

    question_count = len(load_questions(question_file)[:limit] if limit else load_questions(question_file))
    base = {"model": config.CHAT_MODEL, "chunk": config.CHUNK_CHARS, "top_k": config.RETRIEVAL_TOP_K,
            "method": config.RAG_METHOD}
    variants = _variants(base, chunk_sizes, top_ks, models, methods)
    # Aynı modeli kullanan denemeler arka arkaya: model değiştirmek (bellekten çıkarıp yüklemek) zaman alır.
    variants = variants[:1] + sorted(variants[1:], key=lambda v: v["model"] != base["model"])

    needed = sorted({v["model"] for v in variants} | {judge, config.EMBED_MODEL})
    if not llm.health()["ollama"]:
        raise EvaluationError("Ollama'ya ulaşılamıyor. Ollama uygulamasını açıp tekrar deneyin.")
    installed = set(llm.server_info()["models"])
    missing = [m for m in needed if not llm._has_model(installed, m)]
    if missing:
        if not pull_missing:
            raise EvaluationError("Eksik model(ler): " + ", ".join(missing))
        need_gb = sum(MODEL_SIZES_GB.get(m, 0) for m in missing)
        free_gb = shutil.disk_usage(Path.home()).free / 1024 ** 3
        if need_gb and free_gb < need_gb + 3:
            raise EvaluationError(f"İndirilecek modeller ({', '.join(missing)}) yaklaşık {need_gb:.0f} GB yer istiyor, "
                                  f"diskte {free_gb:.0f} GB boş. Yer açıp tekrar deneyin.")
        for model in missing:
            log(f"⬇️  {model} indiriliyor (yalnızca ilk seferde)…")
            last = [-10]

            def progress(pct, last=last):
                if pct >= last[0] + 10:
                    last[0] = pct - pct % 10
                    log(f"   %{last[0]}")
            llm.pull(model, progress)

    out_dir = output_dir / f"karsilastirma_{datetime.now():%Y-%m-%d_%H%M%S}"
    out_dir.mkdir(parents=True, exist_ok=True)
    temp_root = Path(tempfile.mkdtemp(prefix="karsilastirma-"))
    original = (config.CHAT_MODEL, config.CHUNK_CHARS, config.RETRIEVAL_TOP_K, config.RAG_METHOD,
                config.DATA_DIR, config.DB_PATH, config.TABLES_DIR)
    total_runs = len(variants) * repeats
    log(f"\n🔬 {len(variants)} deneme × {repeats} tekrar = {total_runs} değerlendirme, {question_count} soru, hakem: {judge}")
    for v in variants:
        log(f"   • {v['name']:<16} {v['change']}")

    def evaluate(variant: dict) -> dict:
        summaries, reports, chunk_count, prep, node_count = [], [], 0, [], 0
        for attempt in range(1, repeats + 1):
            log(f"\n━━ {variant['name']} ({variant['change']})" + (f" — tekrar {attempt}/{repeats}" if repeats > 1 else ""))
            _apply(variant, temp_root)
            outcome = run_evaluation(question_file, files_dir, out_dir, judge, limit,
                                     log=lambda *a: None,
                                     label=variant["name"] + (f"-{attempt}" if repeats > 1 else ""))
            summaries.append(outcome["summary"])
            reports.append(outcome["report"])
            chunk_count = outcome["chunk_count"]
            prep.append(outcome["prep_seconds"])
            node_count = outcome["node_count"]
            s = outcome["summary"]
            log(f"   doğruluk {_pct(s['score'])} · algılanan bekleme {s['avg_wait']} sn · "
                f"hazırlık {outcome['prep_seconds']} sn · parça {chunk_count}"
                + (f" · bölüm {node_count}" if variant["method"] == "pageindex" else ""))
            if outcome.get("judge_problems"):
                problems = outcome["judge_problems"]
                log(f"   ⚠️  Hakem {len(problems)} soruyu puanlayamadı (0 puan sayıldı). İlk neden: {problems[0]}")
        return _row(variant, {"chunk_count": chunk_count, "reports": reports, "prep": prep,
                              "node_count": node_count}, summaries)

    try:
        rows = [evaluate(v) for v in variants]
        choice = recommend(rows, question_count)
        tested = {(r["model"], r["chunk"], r["top_k"], r["method"]) for r in rows}
        combo = (choice["model"], choice["chunk"], choice["top_k"], choice["method"])
        if combo not in tested:
            # Ayrı ayrı en iyi olan değerler birlikte de iyi mi? Birleşimi doğrula.
            v = {"name": "onerilen", "factor": None, "change": "Önerilen birleşim (doğrulama)",
                 "model": combo[0], "chunk": combo[1], "top_k": combo[2], "method": combo[3]}
            rows.append(evaluate(v))
    finally:
        (config.CHAT_MODEL, config.CHUNK_CHARS, config.RETRIEVAL_TOP_K, config.RAG_METHOD,
         config.DATA_DIR, config.DB_PATH, config.TABLES_DIR) = original
        shutil.rmtree(temp_root, ignore_errors=True)

    path = out_dir / "karsilastirma.xlsx"
    write_comparison(path, rows, choice, judge, question_file, repeats)

    log("\n📊 Karşılaştırma")
    log(f"   {'Deneme':<16} {'Yöntem':<9} {'Model':<11} {'Parça':>5} {'Top-k':>5}  {'Doğruluk':>8}  "
        f"{'Doküman':>7} {'Bulundu':>7} {'Gösterildi':>10}  {'Bekleme':>8}  {'Toplam':>7}  {'Hazırlık':>8}")
    for r in rows:
        mark = "★" if r["name"] == "onerilen" or (r in choice["winners"].values() and r["name"] != "temel") else " "
        log(f" {mark} {r['name']:<16} {METHOD_LABEL[r['method']]:<9} {r['model']:<11} {r['chunk']:>5} {r['top_k']:>5}  "
            f"{_pct(r['score']):>8}  {_pct(r['score_doc']):>7} {_pct(r['retrieved']):>7} {_pct(r['cited']):>10}  "
            f"{_sec(r['wait']):>6} sn  {_sec(r['total']):>5} sn  {_sec(r['prep']):>6} sn")
    log(f"\n✅ Önerilen ayar:  RAG_METHOD={choice['method']}  CHAT_MODEL={choice['model']}  "
        f"CHUNK_CHARS={choice['chunk']}  RETRIEVAL_TOP_K={choice['top_k']}")
    log(f"📄 Tablo: {path}")
    return {"path": path, "rows": rows, "choice": choice}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(prog="ayar_karsilastir",
                                     description="Değerlendirmeyi her seferinde tek ayar değiştirerek tekrarlar.")
    parser.add_argument("sorular", type=Path, help="Soru dosyası (.xlsx)")
    parser.add_argument("--klasor", type=Path, default=None, help="Dosya klasörü (varsayılan: soru dosyasının klasörü)")
    parser.add_argument("--hakem-model", default=DEFAULT_JUDGE,
                        help=f"Bütün denemelerde kullanılacak hakem (varsayılan: {DEFAULT_JUDGE})")
    parser.add_argument("--parca", type=_floats, default=[600, 1200, 2000], help="Parça boyutları (ör. 600,1200,2000)")
    parser.add_argument("--topk", type=_floats, default=[3, 6, 10], help="Bulunan parça sayıları (ör. 3,6,10)")
    parser.add_argument("--modeller", type=lambda s: [m.strip() for m in s.split(",") if m.strip()],
                        default=["qwen3:8b", "qwen3:14b"], help="Cevap modelleri (ör. qwen3:8b,qwen3:14b)")
    parser.add_argument("--yontemler", type=lambda s: [m.strip().lower() for m in s.split(",") if m.strip()],
                        default=[], help="Doküman arama yöntemleri (ör. vector,pageindex)")
    parser.add_argument("--tekrar", type=int, default=1,
                        help="Her denemeyi kaç kez çalıştırıp ortalamasını al (rastlantısallığı azaltır)")
    parser.add_argument("--sadece", type=int, default=None, help="Sadece ilk N soru (hızlı deneme için)")
    parser.add_argument("--cikti", type=Path, default=ROOT / "degerlendirme" / "raporlar")
    args = parser.parse_args(argv)

    if not args.sorular.exists():
        print(f"✖ Soru dosyası bulunamadı: {args.sorular}")
        return 2
    unknown = [m for m in args.yontemler if m not in METHOD_LABEL]
    if unknown:
        print(f"✖ Bilinmeyen yöntem: {', '.join(unknown)} (geçerli: vector, pageindex)")
        return 2
    temp_dir = tempfile.mkdtemp(prefix="karsilastirma-")
    os.environ["DATA_DIR"] = temp_dir
    os.environ.setdefault("ALLOW_REGISTRATION", "false")

    from .calistir import EvaluationError
    from .sorular import QuestionFileError

    try:
        run_comparison(args.sorular, args.klasor or args.sorular.parent, args.cikti, args.hakem_model,
                       args.parca, args.topk, args.modeller, max(1, args.tekrar), args.sadece,
                       methods=args.yontemler)
    except (EvaluationError, QuestionFileError) as e:
        print(f"\n✖ {e}")
        return 1
    except KeyboardInterrupt:
        print("\nDurduruldu.")
        return 130
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
