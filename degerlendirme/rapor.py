"""
rapor.py
--------
Değerlendirme sonuçlarını Excel raporuna yazar. Üç sayfa:
  Özet      — genel puanlar ve soru türüne göre kırılım
  Sonuçlar  — her soru için bir satır: cevap, karar, kaynak, süre…
  Ayarlar   — hangi model ve ayarlarla çalıştığı, yüklenen dosyalar, bilgisayar bilgisi
"""

import statistics
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .puanlama import VERDICT_CORRECT, VERDICT_ERROR, VERDICT_PARTIAL, VERDICT_UNCLEAR, VERDICT_WRONG
from .sorular import TYPES

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
FILLS = {
    VERDICT_CORRECT: PatternFill("solid", fgColor="D1FAE5"),
    VERDICT_PARTIAL: PatternFill("solid", fgColor="FEF3C7"),
    VERDICT_WRONG: PatternFill("solid", fgColor="FEE2E2"),
    VERDICT_ERROR: PatternFill("solid", fgColor="FECACA"),
    VERDICT_UNCLEAR: PatternFill("solid", fgColor="E5E7EB"),
    "evet": PatternFill("solid", fgColor="D1FAE5"),
    "hayır": PatternFill("solid", fgColor="FEE2E2"),
}

RESULT_COLUMNS = [
    ("No", 6), ("Excel satırı", 8), ("Soru türü", 16), ("Soru", 45), ("Doğru cevap", 30),
    ("Sistemin cevabı", 60), ("Karar", 10), ("Hakemin gerekçesi", 40), ("Sayılar tuttu mu?", 12),
    ("Beklenen yol", 11), ("Seçilen yol", 11), ("Yol doğru mu?", 10),
    ("Beklenen kaynak", 26), ("Aramada bulundu mu?", 12), ("Cevapta gösterildi mi?", 12),
    ("Gösterilen kaynaklar", 34), ("Kullanılan SQL", 45), ("Süre (sn)", 9), ("İlk kelime (sn)", 10), ("Hata", 30),
]


def _ratio(values: List[str]) -> Optional[float]:
    """'evet'/'hayır'/'kısmen (1/2)' listesinden başarı oranı; ilgisiz ('—') satırlar sayılmaz."""
    relevant = [v for v in values if v and v != "—"]
    if not relevant:
        return None
    return sum(1 for v in relevant if v == "evet") / len(relevant)


def _score(verdicts: List[str]) -> Optional[float]:
    """Doğru = 1 puan, Kısmen = yarım puan. Hata ve belirsizler de paydaya dahildir (gizlenmesinler)."""
    if not verdicts:
        return None
    points = sum(1.0 if v == VERDICT_CORRECT else 0.5 if v == VERDICT_PARTIAL else 0.0 for v in verdicts)
    return points / len(verdicts)


def _pct(value: Optional[float]) -> str:
    return "—" if value is None else f"%{value * 100:.0f}"


def _sec(values: List[Optional[float]], how) -> str:
    values = [v for v in values if v is not None]
    return "—" if not values else f"{how(values):.1f}"


def summarize(results: List[dict]) -> dict:
    """Terminalde de gösterilen başlıca sayılar."""
    groups = {kind: [r for r in results if r["kind"] == kind] for kind in TYPES}
    summary = {"total": len(results), "score": _score([r["verdict"] for r in results]), "by_type": {}}
    for kind, items in groups.items():
        if not items:
            continue
        summary["by_type"][kind] = {
            "count": len(items),
            "correct": sum(r["verdict"] == VERDICT_CORRECT for r in items),
            "partial": sum(r["verdict"] == VERDICT_PARTIAL for r in items),
            "wrong": sum(r["verdict"] == VERDICT_WRONG for r in items),
            "other": sum(r["verdict"] in (VERDICT_ERROR, VERDICT_UNCLEAR) for r in items),
            "score": _score([r["verdict"] for r in items]),
            "route": _ratio([r["route_ok"] for r in items]),
            "retrieved": _ratio([r["retrieved_ok"] for r in items]),
            "cited": _ratio([r["cited_ok"] for r in items]),
            "numbers": _ratio([r["numbers_ok"] for r in items]),
            "avg_seconds": _sec([r["seconds"] for r in items], statistics.mean),
            "avg_first": _sec([r["first_token_seconds"] for r in items], statistics.mean),
        }
    seconds = [r["seconds"] for r in results]
    summary["avg_seconds"] = _sec(seconds, statistics.mean)
    summary["median_seconds"] = _sec(seconds, statistics.median)
    summary["max_seconds"] = _sec(seconds, max)
    summary["avg_first"] = _sec([r["first_token_seconds"] for r in results], statistics.mean)
    summary["errors"] = sum(r["verdict"] == VERDICT_ERROR for r in results)
    return summary


def _header(ws, row: int, labels: List[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_report(path: Path, results: List[dict], settings: List[tuple], files: List[dict]) -> dict:
    summary = summarize(results)
    wb = Workbook()

    # --- Özet ---
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = "Cevap Kalitesi Değerlendirme Raporu"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = dict(settings).get("Tarih", "")
    rows = [
        ("Toplam soru", summary["total"]),
        ("Genel doğruluk puanı", _pct(summary["score"])),
        ("Hata ile biten soru", summary["errors"]),
        ("Ortalama süre (sn)", summary["avg_seconds"]),
        ("Ortanca süre (sn)", summary["median_seconds"]),
        ("En uzun süre (sn)", summary["max_seconds"]),
        ("Ortalama ilk kelime süresi (sn)", summary["avg_first"]),
        ("Sohbet modeli", dict(settings).get("Sohbet modeli", "")),
        ("Hakem modeli", dict(settings).get("Hakem modeli", "")),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=value)

    start = 4 + len(rows) + 2
    ws.cell(row=start - 1, column=1, value="Soru türüne göre").font = BOLD
    labels = ["Soru türü", "Soru sayısı", "Doğru", "Kısmen", "Yanlış", "Hata/Belirsiz", "Doğruluk puanı",
              "Yol doğru", "Aramada bulundu", "Cevapta gösterildi", "Sayılar tuttu", "Ort. süre (sn)",
              "Ort. ilk kelime (sn)"]
    _header(ws, start, labels)
    for offset, (kind, s) in enumerate(summary["by_type"].items(), start=1):
        values = [kind, s["count"], s["correct"], s["partial"], s["wrong"], s["other"], _pct(s["score"]),
                  _pct(s["route"]), _pct(s["retrieved"]), _pct(s["cited"]), _pct(s["numbers"]),
                  s["avg_seconds"], s["avg_first"]]
        for col, value in enumerate(values, start=1):
            ws.cell(row=start + offset, column=col, value=value)

    notes_row = start + len(summary["by_type"]) + 3
    notes = [
        "Nasıl okunur?",
        "• Doğruluk puanı: Doğru = 1, Kısmen = ½, Yanlış/Hata = 0 puan; soru sayısına bölünür.",
        "• 'dokümanda olmayan' sorularda 'Doğru', sistemin uydurmadan 'bulamadım' dediği anlamına gelir.",
        "• Yol doğru: doküman sorusu dokümanlara, Excel sorusu SQL'e yönlendirildi mi?",
        "• Aramada bulundu: doğru dosya/sayfa, modele verilen parçalar arasında var mıydı? (arama kalitesi)",
        "• Cevapta gösterildi: cevaptaki [n] kaynakları arasında doğru dosya/sayfa var mı? Excel'de: SQL "
        "doğru tabloyu kullandı mı?",
        "• Sayılar tuttu: beklenen cevaptaki her sayı sistemin cevabında/tablosunda geçiyor mu? (modelden "
        "bağımsız, kurallı kontrol)",
        "• Hakem de bir yapay zekâ modelidir ve yanılabilir. Şüpheli satırları 'Sonuçlar' sayfasında gözle "
        "kontrol edin. Hakem, sohbet modeliyle aynıysa kendi cevaplarına karşı hoşgörülü olabilir.",
    ]
    for i, note in enumerate(notes):
        cell = ws.cell(row=notes_row + i, column=1, value=note)
        if i == 0:
            cell.font = BOLD
    ws.column_dimensions["A"].width = 34
    for col in range(2, len(labels) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    # --- Sonuçlar ---
    rs = wb.create_sheet("Sonuçlar")
    _header(rs, 1, [name for name, _ in RESULT_COLUMNS])
    for i, r in enumerate(results, start=1):
        values = [
            i, r["row"], r["kind"], r["question"], r["expected"], r["answer"], r["verdict"], r["reason"],
            r["numbers_ok"], r["expected_route"], r["actual_route"], r["route_ok"], r["expected_source"],
            r["retrieved_ok"], r["cited_ok"], r["shown_sources"], r["sql"], r["seconds"],
            r["first_token_seconds"], r["error"],
        ]
        for col, value in enumerate(values, start=1):
            cell = rs.cell(row=i + 1, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, str) and value in FILLS and col in (7, 9, 12, 14, 15):
                cell.fill = FILLS[value]
    for col, (_, width) in enumerate(RESULT_COLUMNS, start=1):
        rs.column_dimensions[get_column_letter(col)].width = width
    rs.freeze_panes = "E2"
    rs.auto_filter.ref = f"A1:{get_column_letter(len(RESULT_COLUMNS))}{len(results) + 1}"

    # --- Ayarlar ---
    st = wb.create_sheet("Ayarlar")
    _header(st, 1, ["Ayar", "Değer"])
    for i, (key, value) in enumerate(settings, start=2):
        st.cell(row=i, column=1, value=key).font = BOLD
        st.cell(row=i, column=2, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    file_row = len(settings) + 4
    st.cell(row=file_row - 1, column=1, value="Yüklenen dosyalar").font = BOLD
    _header(st, file_row, ["Dosya", "Tür", "Durum", "Ayrıntı"])
    for i, f in enumerate(files, start=1):
        for col, value in enumerate((f["name"], f["type"], f["status"], f["detail"]), start=1):
            st.cell(row=file_row + i, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    st.column_dimensions["A"].width = 34
    st.column_dimensions["B"].width = 60
    st.column_dimensions["C"].width = 12
    st.column_dimensions["D"].width = 50

    wb.save(path)
    return summary
