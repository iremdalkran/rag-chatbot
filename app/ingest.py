"""
ingest.py
---------
Yüklenen dosyadan metni çıkarır ve aramaya uygun parçalara (chunk) böler.

Desteklenen türler: PDF, Word (.docx), düz metin (.txt, .md).
PDF'lerde her parçanın hangi SAYFADAN geldiği saklanır — cevaplarda kaynak gösterebilmek için.
Parçalar cümle sınırlarından bölünür ve komşu parçalar biraz üst üste biner, böylece bir
cümle iki parçanın arasında kopmaz.
"""

import io
import re
from dataclasses import dataclass
from typing import List, Optional

from app import config

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class IngestError(Exception):
    """Kullanıcıya gösterilecek, anlaşılır mesajlı dosya okuma hatası."""


@dataclass
class Page:
    number: Optional[int]  # PDF sayfa numarası; sayfası olmayan türlerde None
    text: str


@dataclass
class Chunk:
    index: int
    page: Optional[int]
    text: str


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    # Satır sonunda kelimenin tire ile bölünmesini birleştir: "doku-\nman" -> "dokuman"
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> List[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise IngestError("Bu PDF şifreli. Lütfen şifresiz bir kopyasını yükleyin.") from e
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            pages.append(Page(number, _clean(page.extract_text() or "")))
    except IngestError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise IngestError("PDF okunamadı. Dosya bozuk olabilir.") from e
    return pages


def _read_docx(data: bytes) -> List[Page]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise IngestError("Word dosyası okunamadı. Dosya bozuk olabilir (eski .doc biçimi desteklenmez).") from e
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [Page(None, _clean("\n\n".join(parts)))]


def _read_text(data: bytes) -> List[Page]:
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return [Page(None, _clean(data.decode(encoding)))]
        except UnicodeDecodeError:
            continue
    raise IngestError("Metin dosyasının karakter kodlaması anlaşılamadı.")


def extract_pages(filename: str, data: bytes) -> List[Page]:
    name = filename.lower()
    if name.endswith(".pdf"):
        pages = _read_pdf(data)
    elif name.endswith(".docx"):
        pages = _read_docx(data)
    elif name.endswith((".txt", ".md")):
        pages = _read_text(data)
    else:
        raise IngestError("Desteklenmeyen dosya türü. PDF, Word (.docx), TXT veya MD yükleyebilirsiniz.")

    if sum(len(p.text) for p in pages) < 20:
        if name.endswith(".pdf"):
            raise IngestError(
                "Bu PDF'ten metin çıkarılamadı. Büyük ihtimalle taranmış (resim olarak kaydedilmiş) bir "
                "belge. Metin seçilebilen bir PDF yükleyin ya da belgeyi önce bir OCR programından geçirin."
            )
        raise IngestError("Dosyada okunabilir metin bulunamadı.")
    return pages


_SENTENCE_SPLIT = re.compile(r"(?<=[.!?…:;])\s+|\n+")


def _split_units(text: str, max_chars: int) -> List[str]:
    """Metni cümlelere böler; tek başına çok uzun olan cümleleri de kelime sınırından keser."""
    units = []
    for sentence in _SENTENCE_SPLIT.split(text):
        sentence = sentence.strip()
        if not sentence:
            continue
        while len(sentence) > max_chars:
            cut = sentence.rfind(" ", 0, max_chars)
            if cut <= 0:
                cut = max_chars
            units.append(sentence[:cut].strip())
            sentence = sentence[cut:].strip()
        if sentence:
            units.append(sentence)
    return units


def chunk_pages(pages: List[Page], size: int = None, overlap: int = None) -> List[Chunk]:
    size = size or config.CHUNK_CHARS
    overlap = config.CHUNK_OVERLAP_CHARS if overlap is None else overlap
    overlap = min(overlap, size // 2)

    chunks: List[Chunk] = []
    for page in pages:
        units = _split_units(page.text, size)
        current: List[str] = []
        length = 0
        for unit in units:
            if current and length + len(unit) + 1 > size:
                chunks.append(Chunk(len(chunks), page.number, " ".join(current)))
                # Bir sonraki parçaya, öncekinin son cümlelerini (overlap kadar) taşı.
                carried: List[str] = []
                carried_len = 0
                for previous in reversed(current):
                    if carried_len + len(previous) + 1 > overlap:
                        break
                    carried.insert(0, previous)
                    carried_len += len(previous) + 1
                current, length = carried, carried_len
            current.append(unit)
            length += len(unit) + 1
        if current:
            chunks.append(Chunk(len(chunks), page.number, " ".join(current)))
    return chunks


def pdf_outline(data: bytes) -> List[dict]:
    """PDF'in kendi içindekiler (yer imi) listesi: [{"level", "title", "page"}]. Yoksa boş liste.
    PageIndex ağacı kurarken önce buna bakar; varsa modelin başlık çıkarmasına gerek kalmaz."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        result: List[dict] = []

        def walk(items, level):
            for item in items:
                if isinstance(item, list):
                    walk(item, level + 1)
                    continue
                try:
                    page = reader.get_destination_page_number(item) + 1
                except Exception:
                    continue
                title = str(getattr(item, "title", "") or "").strip()
                if title and page >= 1:
                    result.append({"level": level, "title": title[:200], "page": page})

        walk(reader.outline, 1)
        return result
    except Exception:
        return []
