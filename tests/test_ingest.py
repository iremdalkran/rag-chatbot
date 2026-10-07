import io
from pathlib import Path

import docx
import pytest
from pypdf import PdfWriter

from app import ingest

SAMPLE_PDF = Path(__file__).resolve().parent.parent / "ornekler" / "ornek_dokuman.pdf"


def test_pdf_pages_are_numbered():
    pages = ingest.extract_pages("ornek.pdf", SAMPLE_PDF.read_bytes())
    assert pages[0].number == 1
    assert "Kahve" in pages[0].text


def test_scanned_pdf_gives_clear_error():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ingest.IngestError, match="taranmış"):
        ingest.extract_pages("bos.pdf", buf.getvalue())


def test_corrupt_pdf_gives_clear_error():
    with pytest.raises(ingest.IngestError):
        ingest.extract_pages("bozuk.pdf", b"%PDF-1.4 bu bir pdf degil")


def test_docx_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Yıllık izin hakkı on dört gündür.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kıdem"
    table.rows[0].cells[1].text = "5 yıl"
    buf = io.BytesIO()
    document.save(buf)
    pages = ingest.extract_pages("politika.docx", buf.getvalue())
    assert "on dört" in pages[0].text and "Kıdem | 5 yıl" in pages[0].text


def test_text_in_windows_turkish_encoding():
    pages = ingest.extract_pages("not.txt", "Şirket çalışanları için ağ kuralları.".encode("cp1254"))
    assert pages[0].text.startswith("Şirket")


def test_unsupported_type():
    with pytest.raises(ingest.IngestError):
        ingest.extract_pages("resim.png", b"123")


def test_chunks_respect_size_overlap_and_pages():
    sentence = "Bu bir deneme cümlesidir ve biraz uzundur."
    pages = [ingest.Page(1, " ".join([sentence] * 60)), ingest.Page(2, "İkinci sayfa kısa.")]
    chunks = ingest.chunk_pages(pages, size=300, overlap=80)
    assert all(len(c.text) <= 300 for c in chunks)
    assert chunks[-1].page == 2 and chunks[-1].text == "İkinci sayfa kısa."
    first, second = chunks[0].text, chunks[1].text
    assert second.startswith(first.split(". ")[-1][:20])  # komşu parçalar üst üste biniyor
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_very_long_word_does_not_loop_forever():
    chunks = ingest.chunk_pages([ingest.Page(None, "a" * 5000)], size=1000, overlap=200)
    assert sum(len(c.text) for c in chunks) == 5000
