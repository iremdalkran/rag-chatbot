"""Değerlendirme ARACININ kendi testleri (cevap kalitesini değil, aracın doğru çalıştığını ölçer)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from degerlendirme import puanlama
from degerlendirme.calistir import run
from degerlendirme.sorular import (TYPE_DATA, TYPE_DOC, TYPE_NONE, Question, QuestionFileError,
                                   load_questions, normalize_type, parse_pages, write_question_file)

EXAMPLES = Path(__file__).resolve().parent.parent / "ornekler" / "degerlendirme"


def q(kind=TYPE_DOC, expected="14 iş günü", file="personel_yonetmeligi.pdf", pages=None):
    return Question(2, "soru?", expected, file, pages, kind)


@pytest.mark.parametrize("value,kind", [
    ("doküman", TYPE_DOC), ("Dokuman", TYPE_DOC), ("PDF", TYPE_DOC), ("excel", TYPE_DATA), ("Veri", TYPE_DATA),
    ("dokümanda olmayan", TYPE_NONE), ("Dokumanda Olmayan", TYPE_NONE), ("", None), ("başka", None),
])
def test_normalize_type(value, kind):
    assert normalize_type(value) == kind


def test_parse_pages():
    assert parse_pages(4) == {4} and parse_pages(4.0) == {4}
    assert parse_pages("4-6") == {4, 5, 6} and parse_pages("2, 7") == {2, 7}
    assert parse_pages("") is None and parse_pages(None) is None


def test_number_check_handles_turkish_formatting():
    assert puanlama.check_numbers("3.500 TL", "Üst sınır 3500 TL'dir [1].") == "evet"
    assert puanlama.check_numbers("2.807.800 TL", "Toplam 2.807.800 TL.") == "evet"
    assert puanlama.check_numbers("12,5", "oran 12.5") == "evet"
    assert puanlama.check_numbers("İstanbul (1.067.200 TL)", "İstanbul",
                                  {"columns": ["b", "t"], "rows": [["İstanbul", 1067200]]}) == "evet"
    assert puanlama.check_numbers("14 iş günü", "20 gün") == "hayır"
    assert puanlama.check_numbers("3 gün, 5 gün", "3 gün") == "kısmen (1/2)"
    assert puanlama.check_numbers("Brezilya", "Brezilya") == "—"


def test_document_source_checks_separate_retrieval_and_citation():
    sources = [{"filename": "kahve.pdf", "page": 1, "cited": True},
               {"filename": "personel_yonetmeligi.pdf", "page": 2, "cited": False}]
    assert puanlama.check_document_sources(q(pages={2}), sources) == ("evet", "hayır")
    assert puanlama.check_document_sources(q(pages={3}), sources) == ("hayır", "hayır")
    sources[1]["cited"] = True
    assert puanlama.check_document_sources(q(pages={2}), sources) == ("evet", "evet")
    assert puanlama.check_document_sources(q(file="personel_yonetmeligi"), sources) == ("evet", "evet")
    assert puanlama.check_document_sources(q(kind=TYPE_NONE, file=""), sources) == ("—", "—")


def test_table_source_check():
    tables = {"satislar.xlsx": ["satislar"]}
    question = q(kind=TYPE_DATA, file="satislar.xlsx")
    assert puanlama.check_table_source(question, 'SELECT SUM(tutar) FROM "satislar"', tables) == "evet"
    assert puanlama.check_table_source(question, "SELECT * FROM satislar_eski", tables) == "hayır"
    assert puanlama.check_table_source(question, None, tables) == "hayır"


def test_question_file_validation(tmp_path):
    path = tmp_path / "s.xlsx"
    write_question_file(path, [("Soru 1", "", "a.pdf", "", "doküman"), ("Soru 2", "x", "", "", "bilinmeyen")])
    with pytest.raises(QuestionFileError) as e:
        load_questions(path)
    assert "2. satır" in str(e.value) and "3. satır" in str(e.value)


def test_template_has_columns_and_dropdown():
    wb = load_workbook(Path(__file__).resolve().parent.parent / "degerlendirme" / "sablon.xlsx")
    ws = wb["Sorular"]
    assert [c.value for c in ws[1]] == ["Soru", "Doğru cevap", "Kaynak dosya", "Kaynak sayfa", "Soru türü"]
    assert ws.max_row == 1  # şablon boş: yanlışlıkla örnek soru sorulmasın
    assert "doküman,Excel,dokümanda olmayan" in ws.data_validations.dataValidation[0].formula1
    assert "Nasıl doldurulur" in wb.sheetnames


def test_full_run_on_example_set_writes_report(ollama, tmp_path):
    report = run(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, log=lambda *_: None)
    wb = load_workbook(report)
    assert wb.sheetnames == ["Özet", "Sonuçlar", "Ayarlar"]

    results = wb["Sonuçlar"]
    header = [c.value for c in results[1]]
    rows = [dict(zip(header, [c.value for c in row])) for row in results.iter_rows(min_row=2)]
    assert len(rows) == 24
    assert {r["Karar"] for r in rows} == {"Doğru"}
    by_question = {r["Soru"]: r for r in rows}
    total = by_question["Toplam satış tutarı ne kadar?"]
    assert total["Seçilen yol"] == "SQL" and total["Yol doğru mu?"] == "evet"
    assert total["Cevapta gösterildi mi?"] == "evet" and "FROM satislar" in total["Kullanılan SQL"]
    leave = by_question["Beş yıldan az çalışan bir personelin yıllık izni kaç gündür?"]
    assert leave["Seçilen yol"] == "Doküman" and leave["Aramada bulundu mu?"] in ("evet", "hayır")
    assert leave["Beklenen kaynak"] == "personel_yonetmeligi.pdf (s. 2)"
    assert all(isinstance(r["Süre (sn)"], (int, float)) for r in rows)
    assert by_question["Şirketin 2025 yılı net kârı ne kadar?"]["Yol doğru mu?"] == "—"

    settings = {row[0].value: row[1].value for row in wb["Ayarlar"].iter_rows(min_row=2) if row[0].value}
    assert settings["Sohbet modeli"].startswith("qwen3:14b (boyut: 14.8B, nicemleme: Q4_K_M")
    assert "AYNI" in settings["Hakem modeli"]
    assert settings["Ollama sürümü"] == "0.40.0"
    assert settings["Aranan parça sayısı (RETRIEVAL_TOP_K)"] == 6
    files = [row[0].value for row in wb["Ayarlar"].iter_rows() if row[0].value and str(row[0].value).endswith((".pdf", ".xlsx"))]
    assert set(files) >= {"kahve_tarihi.pdf", "personel_yonetmeligi.pdf", "satislar.xlsx"}
    assert "sorular.xlsx" not in files

    summary = {row[0].value: row[1].value for row in wb["Özet"].iter_rows(min_row=4, max_row=12)}
    assert summary["Toplam soru"] == 24 and summary["Genel doğruluk puanı"] == "%100"
    # "dokümanda olmayan" sorularda hakeme beklenen cevabın olmadığı söylenmeli.
    none_prompt = [r for r in ollama.judge_requests if "net kârı" in r["messages"][1]["content"]][0]
    assert "dokümanda olmayan" in none_prompt["messages"][1]["content"]


def test_judge_verdicts_and_score(ollama, tmp_path):
    ollama.judge_verdict = "kısmen"
    report = run(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, limit=2, log=lambda *_: None)
    summary = {row[0].value: row[1].value for row in load_workbook(report)["Özet"].iter_rows(min_row=4, max_row=12)}
    assert summary["Toplam soru"] == 2 and summary["Genel doğruluk puanı"] == "%50"


def test_missing_model_stops_with_clear_message(ollama, tmp_path):
    from degerlendirme.calistir import EvaluationError
    with pytest.raises(EvaluationError, match="ollama pull olmayan-model"):
        run(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, judge_model="olmayan-model", log=lambda *_: None)


def test_all_questions_are_asked_before_any_judging(ollama, tmp_path):
    # Hakem farklı bir modelse Ollama her soruda iki modeli değiştirmek zorunda kalmasın.
    run(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, limit=4, log=lambda *_: None)
    kinds = ["judge" if "tarafsız bir hakemsin" in r["messages"][0]["content"] else "ask"
             for r in ollama.chat_requests]
    first_judge = kinds.index("judge")
    assert "ask" not in kinds[first_judge:] and kinds.count("judge") == 4


def test_separate_judge_model_unloads_chat_model_first(ollama, tmp_path):
    report = run(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, judge_model="qwen3:30b-a3b", limit=3,
                 log=lambda *_: None)
    steps = [step for step, _ in ollama.order]
    unload_at = steps.index("unload")
    assert ollama.order[unload_at] == ("unload", "qwen3:14b")
    assert "judge" not in steps[:unload_at] and "ask" not in steps[unload_at:]
    assert {model for step, model in ollama.order if step == "judge"} == {"qwen3:30b-a3b"}
    settings = {row[0].value: row[1].value for row in load_workbook(report)["Ayarlar"].iter_rows(min_row=2) if row[0].value}
    assert settings["Hakem modeli"] == "qwen3:30b-a3b (boyut: 30.5B, nicemleme: Q4_K_M)"


# --- Ayar karşılaştırması ---

from degerlendirme import karsilastir  # noqa: E402


def test_variants_change_one_setting_at_a_time():
    base = {"model": "qwen3:14b", "chunk": 1200, "top_k": 6}
    variants = karsilastir._variants(base, [600, 1200, 2000], [3, 6, 10], ["qwen3:8b", "qwen3:14b"])
    assert [v["name"] for v in variants] == ["temel", "parca-600", "parca-2000", "topk-3", "topk-10", "model-qwen3-8b"]
    for v in variants[1:]:
        differences = [k for k in base if v[k] != base[k]]
        assert len(differences) == 1


def _row(name, factor, score, wait, **cfg):
    return {"name": name, "factor": factor, "score": score, "wait": wait,
            "model": cfg.get("model", "qwen3:14b"), "chunk": cfg.get("chunk", 1200), "top_k": cfg.get("top_k", 6)}


def test_recommendation_prefers_accuracy_then_speed():
    rows = [
        _row("temel", None, 0.90, 8.0),
        _row("parca-600", "chunk", 0.92, 6.0, chunk=600),      # daha hızlı ama en iyiden bir sorudan fazla geride
        _row("parca-2000", "chunk", 0.97, 12.0, chunk=2000),   # en doğru → doğruluk hızdan önce gelir
        _row("topk-3", "top_k", 0.70, 4.0, top_k=3),           # hızlı ama çok kötü → seçilmez
        _row("model-qwen3-8b", "model", 0.88, 5.0, model="qwen3:8b"),  # bir soruluk fark içinde, daha hızlı
    ]
    choice = karsilastir.recommend(rows, question_count=24)
    assert choice["chunk"] == 2000 and choice["top_k"] == 6 and choice["model"] == "qwen3:8b"
    # Sonuç denemelerin sırasına bağlı olmamalı.
    assert karsilastir.recommend(list(reversed(rows[1:])) + rows[:1], 24)["chunk"] == 2000
    # Doğruluklar bir soru içindeyse en hızlısı kazanır.
    close = [_row("temel", None, 0.90, 8.0), _row("parca-600", "chunk", 0.92, 6.0, chunk=600),
             _row("parca-2000", "chunk", 0.93, 12.0, chunk=2000)]
    assert karsilastir.recommend(close, 24)["chunk"] == 600


def test_comparison_runs_variants_and_verifies_combination(ollama, tmp_path):
    from app import config
    before = (config.CHAT_MODEL, config.CHUNK_CHARS, config.RETRIEVAL_TOP_K)
    result = karsilastir.run_comparison(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, "qwen3:30b-a3b",
                                        chunk_sizes=[600, 1200], top_ks=[6], models=["qwen3:8b", "qwen3:14b"],
                                        limit=3, log=lambda *_: None)
    names = [r["name"] for r in result["rows"]]
    assert names[:3] == ["temel", "parca-600", "model-qwen3-8b"]
    assert "qwen3:8b" in ollama.pulled  # eksik model indirildi
    # Her denemede doğru model kullanıldı ve hakem hep aynıydı.
    asked_models = {m for step, m in ollama.order if step == "ask"}
    assert asked_models == {"qwen3:14b", "qwen3:8b"}
    assert {m for step, m in ollama.order if step == "judge"} == {"qwen3:30b-a3b"}
    # Parça boyutu gerçekten değişti: küçük parçalarla daha çok parça oluşur.
    by_name = {r["name"]: r for r in result["rows"]}
    assert by_name["parca-600"]["chunk_count"] > by_name["temel"]["chunk_count"]
    # Ayarlar eski hâline döndü.
    assert (config.CHAT_MODEL, config.CHUNK_CHARS, config.RETRIEVAL_TOP_K) == before

    wb = load_workbook(result["path"])
    ws = wb["Karşılaştırma"]
    header = [c.value for c in ws[4]]
    assert header[:7] == ["Deneme", "Değişen ayar", "Cevap modeli", "Parça boyutu", "Top-k", "Parça sayısı", "Doğruluk"]
    table_names = [ws.cell(row=r, column=1).value for r in range(5, 5 + len(result["rows"]))]
    assert table_names == names
    assert any(str(c.value or "").startswith("CHAT_MODEL=") for row in ws.iter_rows() for c in row)
    assert len(list(result["path"].parent.glob("rapor_*.xlsx"))) == len(result["rows"])


def test_comparison_stops_when_disk_is_too_small(ollama, tmp_path, monkeypatch):
    from collections import namedtuple
    from degerlendirme.calistir import EvaluationError
    usage = namedtuple("usage", "total used free")
    monkeypatch.setattr(karsilastir.shutil, "disk_usage", lambda p: usage(0, 0, 2 * 1024 ** 3))
    with pytest.raises(EvaluationError, match="GB yer istiyor"):
        karsilastir.run_comparison(EXAMPLES / "sorular.xlsx", EXAMPLES, tmp_path, "qwen3:30b-a3b",
                                   [1200], [6], ["qwen3:8b"], limit=1, log=lambda *_: None)
